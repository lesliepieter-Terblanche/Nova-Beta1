"""Turn documents into clean text (Markdown) for reading, answering and filing in the brain.

  • markitdown (Microsoft) — Word, PowerPoint, Excel (.xlsx/.xls), PDF, Outlook .msg, EPUB, HTML, CSV, JSON, ZIP…
  • docling (IBM, optional, a big download) — much better PDFs: real tables, columns, scanned pages (OCR).
    Settings → Documents → Install docling (it gets its own Python, see nova/extras.py). Also pulls tables out
    of a PDF into Excel.
  • the old built-in readers are the fallback, so nothing breaks if a package is missing.
"""
from __future__ import annotations

import importlib.util
from pathlib import Path

from . import context

MARKITDOWN_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".xls", ".msg", ".epub", ".html", ".htm", ".csv", ".json",
                  ".xml", ".zip", ".ipynb", ".rtf"}
DOCLING_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
PLAIN_EXT = {".txt", ".md", ".log", ".eml", ".ini", ".yaml", ".yml", ".py", ".js", ".css"}
READABLE = MARKITDOWN_EXT | PLAIN_EXT

_md = None
_dl = None


def has(module: str) -> bool:
    try:
        return importlib.util.find_spec(module) is not None
    except (ImportError, ValueError):
        return False


def _cfg() -> dict:
    return dict(((context.cfg or {}).get("documents") or {}))


def docling_available() -> bool:
    from . import extras
    return has("docling") or extras.installed("docling")


def docling_on() -> bool:
    """documents.docling: auto (use it when installed) | on | off."""
    mode = str(_cfg().get("docling", "auto")).lower()
    return mode not in ("off", "false", "no") and docling_available()


def _docling_process(mode: str, p: Path) -> str:
    """Run docling in its own Python (see nova/extras.py) and return what it wrote."""
    import subprocess
    import tempfile

    from . import extras
    out = Path(tempfile.mkstemp(suffix=".md" if mode == "md" else ".json")[1])
    try:
        r = subprocess.run([str(extras.python("docling")), str(Path(__file__).with_name("docling_runner.py")), mode,
                            str(p), str(out)], capture_output=True, text=True, timeout=900)
        if r.returncode != 0:
            raise RuntimeError((r.stderr or r.stdout)[-400:])
        return out.read_text(encoding="utf-8")
    finally:
        out.unlink(missing_ok=True)


def _markitdown(p: Path) -> str:
    global _md
    if _md is None:
        from markitdown import MarkItDown
        _md = MarkItDown(enable_plugins=False)
    r = _md.convert(str(p))
    return getattr(r, "markdown", None) or r.text_content or ""


def _docling_converter():
    global _dl
    if _dl is None:
        from docling.document_converter import DocumentConverter
        _dl = DocumentConverter()
    return _dl


def _docling(p: Path) -> str:
    if not has("docling"):
        return _docling_process("md", p)
    return _docling_converter().convert(str(p)).document.export_to_markdown()


def _builtin(p: Path, limit: int) -> str:
    """The original readers (pypdf / python-docx / openpyxl / plain text)."""
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)[:limit]
    if ext == ".docx":
        import docx
        return "\n".join(par.text for par in docx.Document(str(p)).paragraphs)[:limit]
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            rows.append(f"## {ws.title}")
            for row in ws.iter_rows(values_only=True, max_row=300):
                rows.append(", ".join("" if v is None else str(v) for v in row))
        return "\n".join(rows)[:limit]
    return p.read_text(encoding="utf-8", errors="ignore")[:limit]


def to_markdown(path: str | Path, limit: int = 60000, engine: str = "auto") -> str:
    """Best available text for a document. engine: auto | docling | markitdown | builtin."""
    p = Path(path)
    ext = p.suffix.lower()
    order = []
    if engine in ("auto", "docling") and ext in DOCLING_EXT and (engine == "docling" or docling_on()):
        order.append(_docling)
    if engine in ("auto", "markitdown") and ext in MARKITDOWN_EXT and has("markitdown"):
        order.append(_markitdown)
    last = None
    for fn in order:
        try:
            text = fn(p)
            if text and text.strip():
                return text[:limit]
        except Exception as e:                      # a broken file or a missing optional part — try the next one
            last = e
            print(f"[convert] {fn.__name__.strip('_')} couldn't read {p.name}: {e}")
    if ext in (".pdf", ".docx", ".xlsx") or ext in PLAIN_EXT or not order:
        return _builtin(p, limit)
    raise RuntimeError(f"couldn't read {p.name}: {last}")


def pdf_tables(path: str | Path) -> list[dict]:
    """Every table in a document: [{"columns": [...], "rows": [[...], ...]}] (needs docling)."""
    if not docling_available():
        raise RuntimeError("Pulling tables out needs docling — Settings → Documents → Install docling.")
    if not has("docling"):
        import json
        return json.loads(_docling_process("tables", Path(path)))
    doc = _docling_converter().convert(str(path)).document
    out = []
    for t in doc.tables:
        try:
            df = t.export_to_dataframe(doc=doc)
        except TypeError:                           # older docling
            df = t.export_to_dataframe()
        out.append({"columns": [str(c) for c in df.columns],
                    "rows": [["" if v is None else str(v) for v in row] for row in df.itertuples(index=False)]})
    return out
