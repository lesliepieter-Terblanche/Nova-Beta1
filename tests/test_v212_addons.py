"""v2.12 add-ons: markitdown/docling, Silero VAD, ActivityWatch, ComfyUI, moviepy video tools, browser-use web agent.
Everything external (web APIs, subprocesses, models) is faked."""
import datetime as dt
import json
import time
from pathlib import Path

import numpy as np
import pytest

from nova import context
from nova.tools import REGISTRY, select_tools


# ── documents ─────────────────────────────────────────────
def _pptx(path):
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Juniper QBR"
    s.placeholders[1].text = "Pipeline R4m"
    p.save(path)
    return path


def test_markitdown_reads_powerpoint_and_word(nova):
    pytest.importorskip("markitdown")
    tmp = nova[1]
    f = _pptx(tmp / "files" / "qbr.pptx")
    out = REGISTRY["read_file"].run({"path": str(f)})
    assert "Juniper QBR" in out and "Pipeline R4m" in out
    import docx
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "SKU", "Price", "EX2300", "12 000"
    d.save(tmp / "files" / "prices.docx")
    md = REGISTRY["read_file"].run({"path": str(tmp / "files" / "prices.docx")})
    assert "| EX2300 | 12 000 |" in md                     # tables survive as Markdown tables
    res = REGISTRY["document_to_markdown"].run({"path": str(f)})
    assert "Saved" in res and Path(res.split("Saved ")[1].split(" (")[0]).read_text().count("Juniper QBR") == 1


def test_converter_falls_back_and_prefers_docling(nova, monkeypatch):
    from nova import convert
    f = nova[1] / "files" / "a.pdf"
    f.write_bytes(b"%PDF-1.4 not really")
    monkeypatch.setattr(convert, "docling_available", lambda: True)
    monkeypatch.setattr(convert, "_docling", lambda p: "# From docling\n| a | b |")
    assert convert.to_markdown(f).startswith("# From docling")
    nova[0]["documents"] = {"docling": "off"}
    monkeypatch.setattr(convert, "_markitdown", lambda p: "from markitdown")
    assert convert.to_markdown(f) == "from markitdown"
    monkeypatch.setattr(convert, "has", lambda m: False)                   # nothing installed: old readers
    t = nova[1] / "files" / "n.txt"
    t.write_text("plain notes")
    assert convert.to_markdown(t) == "plain notes"


def test_extract_tables_to_excel_with_real_numbers(nova, monkeypatch):
    import openpyxl

    from nova import convert
    f = nova[1] / "files" / "price list.pdf"
    f.write_bytes(b"%PDF")
    monkeypatch.setattr(convert, "docling_available", lambda: True)
    monkeypatch.setattr(convert, "_docling_process", lambda mode, p: json.dumps(
        [{"columns": ["SKU", "Price", "Note"], "rows": [["EX2300", "12 000", "1,5"], ["MX204", "1,299.50", "new"]]}]))
    monkeypatch.setattr(convert, "has", lambda m: m != "docling")
    out = REGISTRY["extract_tables"].run({"path": str(f)})
    assert "Found 1 table(s) (2×3)" in out
    ws = openpyxl.load_workbook(out.split("saved to ")[1].rstrip(".")).active
    assert [c.value for c in ws[2]] == ["EX2300", 12000, "1,5"] and ws["B3"].value == 1299.5


def test_extract_tables_says_how_to_get_docling(nova, monkeypatch):
    from nova import convert
    monkeypatch.setattr(convert, "docling_available", lambda: False)
    f = nova[1] / "files" / "x.pdf"
    f.write_bytes(b"%PDF")
    assert "Install docling" in REGISTRY["extract_tables"].run({"path": str(f)})


def test_inbox_files_powerpoints():
    from nova.inbox import TEXT_EXT
    assert {".pptx", ".xls", ".msg", ".epub"} <= TEXT_EXT


# ── Silero VAD ────────────────────────────────────────────
class FakeSession:
    """Says 'speech' when the chunk is loud; checks the v5 input layout (64 context + 512 samples)."""
    def __init__(self):
        self.calls = []

    def run(self, _, feeds):
        x = feeds["input"]
        assert x.shape == (1, 576) and feeds["state"].shape == (2, 1, 128)
        self.calls.append(x)
        loud = float(np.abs(x[0, 64:]).mean()) > 0.05
        return [np.array([[0.9 if loud else 0.02]], dtype=np.float32), feeds["state"]]


def _vad():
    from nova.vad import SileroVAD
    v = SileroVAD.__new__(SileroVAD)
    v.session, v.threshold = FakeSession(), 0.5
    v.reset()
    return v


def test_vad_chunks_blocks_and_carries_leftovers():
    v = _vad()
    quiet = np.zeros(1280, np.int16)
    assert v.speech_in(quiet) < 0.5 and len(v.session.calls) == 2 and len(v._carry) == 256
    loud = (np.sin(np.arange(1280) / 3) * 12000).astype(np.int16)
    assert v.is_speech(loud) and len(v.session.calls) == 5     # 256 + 1280 = 3 chunks of 512
    assert np.allclose(v._context, v.session.calls[-1][:, -64:])


def test_voice_loop_records_by_speech_not_loudness(nova):
    from nova.voice import BLOCK, VoiceLoop
    cfg = nova[0]
    loop = VoiceLoop(cfg, agent=None, speech=None)
    loop._vad, loop._vad_tried = _vad(), True
    fan = (np.random.default_rng(1).normal(0, 3, BLOCK)).astype(np.int16)        # quiet hum: not speech
    talk = (np.sin(np.arange(BLOCK) / 3) * 12000).astype(np.int16)

    class Stream:
        def __init__(self, blocks):
            self.blocks = list(blocks)

        def read(self, n):
            b = self.blocks.pop(0) if self.blocks else fan
            return (b.reshape(-1, 1), False)
    assert loop._record(Stream([fan] * 80), start_timeout=1.0) is None
    audio = loop._record(Stream([fan] * 3 + [talk] * 15 + [fan] * 40), start_timeout=2.0)
    assert audio is not None and 1.2 < len(audio) / 16000 < 3.5          # stops ~1.1 s after the talking


def test_vad_falls_back_to_loudness_when_unavailable(nova, monkeypatch):
    from nova import vad
    from nova.voice import VoiceLoop
    monkeypatch.setattr(vad, "SileroVAD", lambda **k: (_ for _ in ()).throw(RuntimeError("no model")))
    assert VoiceLoop(nova[0], None, None).vad() is None
    nova[0]["voice"]["vad"] = "loudness"
    assert VoiceLoop(nova[0], None, None).vad() is None


# ── ActivityWatch ─────────────────────────────────────────
class Resp:
    def __init__(self, data=None, status=200, content=b""):
        self._d, self.status_code, self.content = data, status, content
        self.text = json.dumps(data) if data is not None else ""

    def json(self):
        return self._d

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(self.status_code)


AW_EVENTS = [
    {"duration": 5400, "data": {"app": "EXCEL.EXE", "title": "Juniper forecast.xlsx - Excel"}},
    {"duration": 1500, "data": {"app": "chrome.exe", "title": "Funny cats - YouTube - Google Chrome"}},
    {"duration": 900, "data": {"app": "OUTLOOK.EXE", "title": "Inbox - Outlook"}},
    {"duration": 300, "data": {"app": "chrome.exe", "title": "Facebook - Google Chrome"}},
]


@pytest.fixture()
def aw(monkeypatch):
    from nova import activity
    seen = []

    def post(url, json=None, timeout=None):
        seen.append(json)
        if "aw-watcher-web" in " ".join(json["query"]):
            return Resp([[{"duration": 1400, "data": {"url": "https://www.youtube.com/watch?v=1", "title": "Cats"}}]])
        return Resp([AW_EVENTS])
    monkeypatch.setattr(activity.httpx, "post", post)
    monkeypatch.setattr(activity.httpx, "get", lambda url, timeout=None: Resp({"hostname": "pc"}))
    activity._last_nudge = 0.0
    return activity, seen


def test_screen_time_summary_and_tools(nova, aw):
    activity, seen = aw
    s = activity.summary("today")
    assert s["total_text"] == "2 h 15 min" and s["apps"][0] == {"app": "Excel", "seconds": 5400, "text": "1 h 30 min",
                                                                "share": 67}
    assert s["distracted_text"] == "30 min" and s["drift"][0]["what"] == "youtube"
    q = " ".join(seen[0]["query"])
    assert "not-afk" in q and "/" in seen[0]["timeperiods"][0]                 # away time is left out
    out = REGISTRY["screen_time"].run({})
    assert "Active screen time today: 2 h 15 min" in out and "• Excel: 1 h 30 min (67%)" in out and "Drift" in out
    assert "25 min on youtube today" in REGISTRY["time_spent_on"].run({"what": "youtube"})
    assert "screen_time" in {t.name for t in select_tools("where did my day go?")}


def test_screen_time_when_activitywatch_is_missing(nova, monkeypatch):
    from nova import activity

    def boom(*a, **k):
        raise ConnectionError("refused")
    monkeypatch.setattr(activity.httpx, "get", boom)
    assert "activitywatch.net" in REGISTRY["screen_time"].run({})


def test_drift_nudge_only_in_work_hours_with_a_focus_task(nova, aw, monkeypatch):
    activity, _ = aw
    from nova import focus
    said = []
    monkeypatch.setattr(context, "announce", said.append)
    tue = dt.datetime(2026, 9, 29, 10, 30)
    assert activity.check_drift(tue, clock=10_000) is None                 # no focus task → leave him be
    focus.add("Send Axiz the Mist pricing", "now")
    msg = activity.check_drift(tue, clock=10_000)
    assert msg and "Youtube" in msg and "Send Axiz the Mist pricing" in msg and said
    assert activity.check_drift(tue, clock=10_600) is None                 # not again within 45 min
    assert activity.check_drift(tue, clock=13_000) is not None
    activity._last_nudge = 0
    assert activity.check_drift(dt.datetime(2026, 9, 27, 10, 30), clock=99_999) is None   # Sunday
    assert activity.check_drift(dt.datetime(2026, 9, 29, 19, 0), clock=99_999) is None    # after work
    nova[0]["activitywatch"]["nudges"] = False
    assert activity.check_drift(tue, clock=99_999) is None


# ── ComfyUI ───────────────────────────────────────────────
PNG = b"\x89PNG\r\n\x1a\nfake"


@pytest.fixture()
def comfy_server(monkeypatch):
    from nova import comfy
    calls = {"prompts": []}

    def get(url, params=None, timeout=None):
        if url.endswith("/system_stats"):
            return Resp({"system": {}})
        if "/object_info/CheckpointLoaderSimple" in url:
            return Resp({"CheckpointLoaderSimple": {"input": {"required": {"ckpt_name": [
                ["dreamshaper_8.safetensors", "sdxl_turbo_1.0_fp16.safetensors"], {}]}}}})
        if "/history/" in url:
            calls["polls"] = calls.get("polls", 0) + 1
            if calls["polls"] < 2:
                return Resp({})
            return Resp({"p1": {"status": {"status_str": "success"}, "outputs": {"9": {"images": [
                {"filename": "nova_0001.png", "subfolder": "", "type": "output"},
                {"filename": "nova_0002.png", "subfolder": "", "type": "output"}]}}}})
        if url.endswith("/view"):
            return Resp(content=PNG)
        raise AssertionError(url)

    def post(url, json=None, timeout=None):
        calls["prompts"].append(json["prompt"])
        return Resp({"prompt_id": "p1"})
    monkeypatch.setattr(comfy.httpx, "get", get)
    monkeypatch.setattr(comfy.httpx, "post", post)
    monkeypatch.setattr(comfy.time, "sleep", lambda s: None)
    return comfy, calls


def test_comfyui_generates_images(nova, comfy_server):
    comfy, calls = comfy_server
    out = REGISTRY["generate_image"].run({"prompt": "a lighthouse at dusk", "format": "vertical", "count": 2})
    assert "Made 2 images with ComfyUI" in out
    wf = calls["prompts"][0]
    assert wf["4"]["inputs"]["ckpt_name"] == "dreamshaper_8.safetensors"
    assert (wf["5"]["inputs"]["width"], wf["5"]["inputs"]["height"]) == (432, 768)
    assert wf["6"]["inputs"]["text"] == "a lighthouse at dusk" and wf["3"]["inputs"]["steps"] == 25
    files = sorted((nova[1] / "files").parent.glob("**/img_*.png")) or sorted(Path("workspace/images").glob("img_*"))
    assert files and files[-1].read_bytes() == PNG


def test_comfyui_model_choice_and_own_workflow(nova, comfy_server, tmp_path):
    comfy, calls = comfy_server
    nova[0]["media"]["comfyui"] = {"checkpoint": "turbo"}
    comfy.generate("a cat", "square", out_dir=tmp_path)
    wf = calls["prompts"][-1]
    assert wf["4"]["inputs"]["ckpt_name"].startswith("sdxl_turbo") and wf["3"]["inputs"]["steps"] == 4
    assert wf["3"]["inputs"]["cfg"] == 1.0 and wf["5"]["inputs"]["width"] == 512
    mine = tmp_path / "wf.json"
    mine.write_text(json.dumps({"1": {"inputs": {"text": "{{prompt}}, cinematic", "seed": "{{seed}}",
                                                 "w": "{{width}}"}}}))
    nova[0]["media"]["comfyui"] = {"workflow": str(mine)}
    comfy.generate("a dog", "landscape", seed=7, out_dir=tmp_path)
    wf = calls["prompts"][-1]
    assert wf["1"]["inputs"] == {"text": "a dog, cinematic", "seed": 7, "w": 1024}


def test_no_image_generator_running(nova, monkeypatch):
    from nova import comfy
    monkeypatch.setattr(comfy, "running", lambda: False)
    assert "Start ComfyUI" in REGISTRY["generate_image"].run({"prompt": "x"})


# ── moviepy video tools ───────────────────────────────────
def test_caption_groups_break_at_pauses_and_sentences():
    from nova.skills.video_edit import caption_groups
    words = [(0, .3, "Welcome"), (.3, .6, "to"), (.6, .9, "the"), (.9, 1.2, "QBR."), (1.2, 1.5, "Pipeline"),
             (2.5, 2.8, "grew"), (2.8, 3.0, "fast")]
    assert [g[2] for g in caption_groups(words)] == ["Welcome to the", "QBR.", "Pipeline", "grew fast"]


def test_vertical_frames_blur_and_crop():
    from nova.skills.video_edit import _vertical_frame
    land = np.zeros((360, 640, 3), np.uint8)
    land[:, 300:340] = 255
    for mode in ("blur", "crop"):
        f = _vertical_frame(land, mode)
        assert f.shape == (1920, 1080, 3) and f.dtype == np.uint8
    blur = _vertical_frame(land, "blur")
    assert blur[960, 540].max() == 255 and blur[100, 540].max() < 200     # sharp middle, soft darker background


def test_pick_moments_from_the_model_and_fallback(nova):
    from nova.skills.video_edit import pick_moments
    segs = [(0, 10, "intro"), (10, 50, "the big idea"), (50, 120, "more")]
    context.llm.complete = lambda *a, **k: json.dumps({"clips": [{"start": 10, "end": 48, "title": "Big idea"},
                                                                 {"start": 100, "end": 400, "title": "too long"}]})
    assert pick_moments(segs, 120, 2, 40) == [{"start": 10.0, "end": 48.0, "title": "Big idea"}]
    context.llm.complete = lambda *a, **k: "no idea"
    even = pick_moments(segs, 120, 2, 40)
    assert len(even) == 2 and all(e["end"] - e["start"] == 40 for e in even)


def _tiny_video(path: Path, seconds=2):
    import imageio_ffmpeg
    import subprocess
    subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-loglevel", "error", "-y", "-f", "lavfi", "-i",
                    f"testsrc=size=320x180:rate=15:duration={seconds}", "-f", "lavfi", "-i",
                    f"sine=frequency=440:duration={seconds}", "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p",
                    "-c:a", "aac", str(path)], check=True)
    return path


def test_shorts_are_vertical_and_captioned(nova, monkeypatch):
    pytest.importorskip("moviepy")
    from moviepy import VideoFileClip

    from nova.skills import video_edit
    src = _tiny_video(nova[1] / "files" / "talk.mp4", 8)
    monkeypatch.setattr(video_edit, "_out", lambda name: nova[1] / f"{name}.mp4")
    monkeypatch.setattr(video_edit, "VERTICAL", (270, 480))            # small frames: fast test
    context.speech.transcribe_words = lambda p: [(1.2, 1.5, "Hello"), (1.5, 1.9, "Juniper"), (2.2, 2.6, "partners")]
    context.speech.transcribe_segments = lambda p: [(0.0, 8.0, "Hello Juniper partners")]
    context.llm.complete = lambda *a, **k: json.dumps({"clips": [{"start": 1.0, "end": 7.0, "title": "Hello"}]})
    out = REGISTRY["video_to_shorts"].run({"path": str(src), "count": 1, "seconds": 15})
    assert "Made 1 vertical clip" in out
    made = nova[1] / "Hello.mp4"
    clip = VideoFileClip(str(made))
    assert clip.size == [270, 480] and 5.8 < clip.duration < 6.2
    plain, captioned = clip.get_frame(3.5), clip.get_frame(0.7)       # "HELLO JUNIPER" shows 0.2-0.9 s into the clip
    assert plain.shape == (480, 270, 3)
    band = slice(int(480 * 0.62), int(480 * 0.78))
    assert (captioned[band] > 245).all(axis=2).sum() > (plain[band] > 245).all(axis=2).sum() + 200   # white text
    clip.close()


def test_video_tools_are_offered(nova):
    names = {t.name for t in select_tools("turn my webinar into tiktok shorts with captions")}
    assert {"video_to_shorts", "add_captions"} <= names


# ── browser-use web agent ─────────────────────────────────
def test_web_agent_picks_a_capable_model(nova, monkeypatch):
    from nova import web_agent
    for k in ("GEMINI_API_KEY", "GROQ_API_KEY", "XAI_API_KEY"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("GROQ_API_KEY", "gsk_x")
    assert web_agent.pick_llm() == {"provider": "groq", "model": "", "api_key": "gsk_x"}
    monkeypatch.setenv("GEMINI_API_KEY", "g-key")
    assert web_agent.pick_llm()["provider"] == "gemini" and web_agent.pick_llm()["model"] == "gemini-2.5-flash"
    nova[0]["web_agent"]["provider"] = "ollama"
    assert web_agent.pick_llm()["base_url"] == "http://localhost:11434"


def test_web_agent_needs_installing_first(nova, monkeypatch):
    from nova import extras
    monkeypatch.setattr(extras, "installed", lambda name: False)
    assert "isn't installed yet" in REGISTRY["web_agent"].run({"task": "find flights"})


def test_web_agent_runs_in_its_own_process_and_reports(nova, monkeypatch, tmp_path):
    import subprocess
    import sys

    from nova import extras, web_agent
    monkeypatch.setenv("GEMINI_API_KEY", "g-key")
    monkeypatch.setattr(extras, "installed", lambda name: True)
    monkeypatch.setattr(web_agent, "HOME", tmp_path)
    monkeypatch.setattr(web_agent, "PROFILE", tmp_path / "profile")
    fake = tmp_path / "runner.py"
    fake.write_text(
        "import json, os, sys\n"
        "job = json.load(open(sys.argv[1]))\n"
        "assert os.environ['NOVA_WEB_AGENT_KEY'] == 'g-key' and 'api_key' not in job\n"
        "assert 'ALLOWED: submit' in job['rules'] and job['provider'] == 'gemini'\n"
        "print(json.dumps({'step': 1, 'goal': 'Open the Juniper partner portal', 'url': 'https://x'}), flush=True)\n"
        "print('some browser-use log line')\n"
        "print(json.dumps({'done': True, 'result': 'Deal registered: #4411', 'success': True}), flush=True)\n")
    monkeypatch.setattr(web_agent, "python", lambda: Path(sys.executable))
    real_popen = subprocess.Popen
    monkeypatch.setattr(web_agent.subprocess, "Popen",
                        lambda args, **kw: real_popen([args[0], str(fake), args[2]], **kw))
    got = []
    msg = web_agent.start("register the Axiz deal on the Juniper portal", "submit", notify=got.append)
    assert msg.startswith("On it")
    for _ in range(100):
        if got:
            break
        time.sleep(0.05)
    assert got and "Web agent done." in got[0] and "Deal registered: #4411" in got[0]
    assert web_agent._state["job"]["steps"][0]["goal"] == "Open the Juniper partner portal"
    assert not list(tmp_path.glob("jobs/*.json"))                       # job file cleaned up
    assert "Deal registered" in web_agent.status()


def test_web_agent_rules_default_to_stopping():
    from nova import web_agent
    assert web_agent._allowed("") == "nothing — stop before every final action"
    assert web_agent._allowed("submit; send") == "submit, send"


def test_extras_install_in_their_own_python(nova, monkeypatch, tmp_path):
    from nova import extras
    monkeypatch.setattr(extras, "ROOT", tmp_path)
    ran = []

    def run(args, **kw):
        ran.append(args)
        if "venv" in args:
            py = extras.python("docling")
            py.parent.mkdir(parents=True, exist_ok=True)
            py.write_text("")

        class R:
            returncode, stdout, stderr = 0, "", ""
        return R()
    monkeypatch.setattr(extras.subprocess, "run", run)
    done = []
    extras._state.pop("docling", None)
    assert "Installing Docling" in extras.install("docling", done.append)
    for _ in range(100):
        if done:
            break
        time.sleep(0.02)
    assert done[0]["installed"] and extras.installed("docling")
    assert any("docling>=2.40" in a for a in ran[-1]) and str(tmp_path / "tools" / "docling") in str(ran[0])


def test_settings_checks_for_the_new_addons(nova, monkeypatch):
    from nova import activity, comfy, settings
    monkeypatch.setattr(settings, "load_config", lambda *a: nova[0], raising=False)
    monkeypatch.setattr(comfy, "running", lambda: False)
    monkeypatch.setattr(activity, "running", lambda: False)
    assert "isn't running" in settings.run_test("comfyui")["message"]
    assert "activitywatch.net" in settings.run_test("activitywatch")["message"]
    ids = {s["id"] for s in settings.SCHEMA}
    assert {"documents", "images", "web_agent", "activitywatch"} <= ids
